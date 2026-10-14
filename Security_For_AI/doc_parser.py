"""
doc_parser.py — Document text extractor & hidden content analyzer.

Extracts text from uploaded files (PDF, DOCX, XLSX, TXT, CSV, JSON, Code, etc.)
and detects steganography / hidden zero-width unicode obfuscations.
"""

import io
import re
from dataclasses import dataclass
from typing import Optional, Tuple, List

# Zero-width unicode characters commonly used to hide malicious text or bypass regex filters
_ZERO_WIDTH_CHARS = {
    "\u200b": "Zero-Width Space (U+200B)",
    "\u200c": "Zero-Width Non-Joiner (U+200C)",
    "\u200d": "Zero-Width Joiner (U+200D)",
    "\ufeff": "Byte Order Mark / Zero-Width No-Break (U+FEFF)",
    "\u2060": "Word Joiner (U+2060)",
    "\u200e": "Left-to-Right Mark (U+200E)",
    "\u200f": "Right-to-Left Mark (U+200F)",
}

_ZERO_WIDTH_PATTERN = re.compile(r"[\u200b\u200c\u200d\ufeff\u2060\u200e\u200f]")


@dataclass
class ExtractedDocument:
    filename: str
    file_type: str
    file_size_bytes: int
    text: str
    char_count: int
    hidden_char_count: int
    hidden_char_types: List[str]
    is_supported: bool
    error_msg: Optional[str] = None


def extract_text_from_file(file_bytes: bytes, filename: str) -> ExtractedDocument:
    """
    Extract readable text from file bytes based on file extension.
    Returns an ExtractedDocument object.
    """
    lower_name = filename.lower()
    file_size = len(file_bytes)
    extracted_text = ""
    error_msg = None
    file_type = "unknown"

    try:
        if lower_name.endswith(".pdf"):
            file_type = "PDF Document"
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(file_bytes))
            pages_text = []
            for idx, page in enumerate(reader.pages):
                page_t = page.extract_text() or ""
                if page_t.strip():
                    pages_text.append(f"--- Page {idx+1} ---\n{page_t}")
            extracted_text = "\n\n".join(pages_text)

        elif lower_name.endswith(".docx"):
            file_type = "Word Document (.docx)"
            import docx
            doc = docx.Document(io.BytesIO(file_bytes))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            # Also extract text inside tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs.append(row_text)
            extracted_text = "\n".join(paragraphs)

        elif lower_name.endswith((".xlsx", ".xls")):
            file_type = "Excel Spreadsheet"
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
            sheet_lines = []
            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                sheet_lines.append(f"--- Sheet: {sheet_name} ---")
                for row in sheet.iter_rows(values_only=True):
                    row_vals = [str(v) for v in row if v is not None]
                    if row_vals:
                        sheet_lines.append("\t".join(row_vals))
            extracted_text = "\n".join(sheet_lines)

        elif lower_name.endswith((
            ".txt", ".md", ".csv", ".tsv", ".json", ".yaml", ".yml",
            ".py", ".js", ".html", ".xml", ".sql", ".log", ".env", ".ini", ".sh"
        )):
            file_type = "Plain Text / Code"
            for encoding in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
                try:
                    extracted_text = file_bytes.decode(encoding)
                    break
                except UnicodeDecodeError:
                    continue
        else:
            # Attempt plain text decoding as fallback
            file_type = "Generic Document"
            try:
                extracted_text = file_bytes.decode("utf-8", errors="replace")
            except Exception as e:
                error_msg = f"Unsupported file format: {filename}"

    except Exception as e:
        error_msg = f"Failed to extract text from {filename}: {str(e)}"

    # Check for hidden zero-width unicode obfuscation
    hidden_matches = _ZERO_WIDTH_PATTERN.findall(extracted_text)
    hidden_char_count = len(hidden_matches)
    hidden_char_types = list({_ZERO_WIDTH_CHARS.get(c, "Zero-width character") for c in hidden_matches})

    return ExtractedDocument(
        filename=filename,
        file_type=file_type,
        file_size_bytes=file_size,
        text=extracted_text,
        char_count=len(extracted_text),
        hidden_char_count=hidden_char_count,
        hidden_char_types=hidden_char_types,
        is_supported=(error_msg is None and len(extracted_text.strip()) > 0),
        error_msg=error_msg,
    )
